#!/usr/bin/env python3
"""Generate ATS-friendly resume (DOCX + PDF + TXT) for Pamella Shome."""

from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from reportlab.lib.colors import HexColor, black
from reportlab.lib.enums import TA_CENTER, TA_JUSTIFY, TA_LEFT
from reportlab.lib.pagesizes import letter
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib.units import inch
from reportlab.platypus import HRFlowable, Paragraph, SimpleDocTemplate, Spacer

OUT_DIR = Path(__file__).resolve().parent
NAME = "PAMELLA SHOME"
CONTACT = (
    "Pune, India  |  +91-8918075806  |  Pamellashome5@gmail.com  |  "
    "linkedin.com/in/pamella-shome-a31089238"
)

SUMMARY = (
    "B.Tech graduate with hands-on Python and Java development across AI/ML, "
    "computer vision, REST APIs, automation, and data processing. Strong foundation "
    "in Object Oriented Programming (OOP), Data Structures and Algorithms (DSA), "
    "SQL, Git, and Software Development Life Cycle (SDLC). Collaborative, "
    "customer-centric engineer seeking the AI/ML Computational Science Associate "
    "role to design, test, and maintain applications using Python and emerging AI."
)

SKILL_BLOCKS = [
    (
        "Programming and Core CS",
        "Python, Java, C++, Kotlin, Object Oriented Programming (OOP), "
        "Data Structures and Algorithms (DSA), Software Development Life Cycle (SDLC)",
    ),
    (
        "Backend, Data, and Tools",
        "SQL, MySQL, DBMS, REST APIs, Flask, Git, GitHub, debugging, testing, "
        "technical documentation, Docker",
    ),
    (
        "AI / ML and Analytics",
        "Machine Learning, Deep Learning, Computer Vision, OpenCV, NLP, LLMs, "
        "Vision-Language Models (VLM), data analysis, data processing, reporting",
    ),
    (
        "Platforms",
        "Android Studio, Android Automotive OS, VS Code, MATLAB, PowerShell",
    ),
]

EDAG_BULLETS = [
    "Develop, test, and maintain Java middleware for an Android Automotive See-Through surround-view application that ingests live video from multiple vehicle-mounted cameras.",
    "Optimize the middleware pipeline for real-time use: fisheye-to-spherical projection, multi-camera stitching and blending, seam feathering, and exposure matching for a stable driver surround view.",
    "Implement REST-style HTTP streaming endpoints between camera capture, middleware, and the in-vehicle display; debug video, performance, and connectivity issues.",
    "Follow SDLC, OOP, Git, coding standards, and code reviews; write technical documentation and collaborate with hardware, QA, and application teams.",
]

VLM_BULLETS = [
    "Building a Python application that integrates a Vision-Language Model (VLM) service so camera and image inputs can be queried in natural language for scene understanding with the See-Through pipeline.",
    "Connecting containerized client, middleware, and VLM components over REST APIs, including payloads, configuration, and service health checks.",
    "Diagnosed connectivity failures (connection refused) using curl and lsof; isolated root cause across config, ports, and container networking and documented the fix.",
]

AUTO_BULLETS = [
    "Developed a Python automation script (openpyxl, python-dotenv) to process QA spreadsheets and translate German test cases to English; migrated the backend to the DeepL REST API while preserving markup, images, and placeholders.",
]

RENTAL_BULLETS = [
    "Built a Python analytics tool using web-scraped Gurugram rental data (BeautifulSoup), data cleaning, and regression/boosting models for price prediction; delivered visualizations and client reports.",
]

CANCER_BULLETS = [
    "Developed a Python machine learning pipeline (Random Forest, SVM, XGBoost, LightGBM, CatBoost voting ensemble) with Recursive Feature Elimination (RFE) to predict cervical cancer risk from patient history.",
]


def set_run_font(run, name="Calibri", size=10, bold=False, color=None):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    run.font.size = Pt(size)
    run.bold = bold
    if color:
        run.font.color.rgb = RGBColor(*color)


def add_bottom_border(paragraph):
    p = paragraph._p
    pPr = p.get_or_add_pPr()
    pBdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "10")
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), "1F4E79")
    pBdr.append(bottom)
    pPr.append(pBdr)


def heading(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.line_spacing = 1.0
    run = p.add_run(text.upper())
    set_run_font(run, size=11, bold=True, color=(31, 78, 121))
    add_bottom_border(p)
    return p


def body(doc, text, size=10, space_after=2, space_before=0):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(space_before)
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.line_spacing = 1.05
    run = p.add_run(text)
    set_run_font(run, size=size)
    return p


def role_line(doc, left, right):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(5)
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.0
    tab_stops = p.paragraph_format.tab_stops
    tab_stops.add_tab_stop(Inches(7.1), WD_TAB_ALIGNMENT.RIGHT)
    run = p.add_run(left)
    set_run_font(run, size=10.5, bold=True)
    run3 = p.add_run("\t" + right)
    set_run_font(run3, size=9.5)
    return p


def italic_line(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.line_spacing = 1.0
    sr = p.add_run(text)
    set_run_font(sr, size=9.5)
    sr.italic = True
    return p


def bullet(doc, text):
    p = doc.add_paragraph(style="List Bullet")
    p.clear()
    p.paragraph_format.left_indent = Inches(0.2)
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(1)
    p.paragraph_format.line_spacing = 1.05
    run = p.add_run(text)
    set_run_font(run, size=10)
    return p


def build_docx(path: Path):
    doc = Document()
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.left_margin = Inches(0.65)
    section.right_margin = Inches(0.65)
    section.top_margin = Inches(0.45)
    section.bottom_margin = Inches(0.45)

    styles = doc.styles["Normal"]
    styles.font.name = "Calibri"
    styles.font.size = Pt(10)

    name = doc.add_paragraph()
    name.alignment = WD_ALIGN_PARAGRAPH.CENTER
    name.paragraph_format.space_after = Pt(1)
    name.paragraph_format.space_before = Pt(0)
    r = name.add_run(NAME)
    set_run_font(r, size=16, bold=True, color=(31, 78, 121))

    contact = doc.add_paragraph()
    contact.alignment = WD_ALIGN_PARAGRAPH.CENTER
    contact.paragraph_format.space_after = Pt(1)
    contact.paragraph_format.space_before = Pt(0)
    cr = contact.add_run(CONTACT)
    set_run_font(cr, size=9.5)

    heading(doc, "Professional Summary")
    body(doc, SUMMARY, size=10, space_after=1)

    heading(doc, "Technical Skills")
    for label, items in SKILL_BLOCKS:
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(1)
        p.paragraph_format.line_spacing = 1.05
        r1 = p.add_run(label + ": ")
        set_run_font(r1, size=10, bold=True)
        r2 = p.add_run(items)
        set_run_font(r2, size=10)

    heading(doc, "Professional Experience")
    role_line(doc, "EDAG Production Solutions India Pvt Ltd", "Pune, India  |  Jun 2024 – Present")
    italic_line(
        doc,
        "Android Automotive Developer  |  See-Through Surround-View Application",
    )
    for t in EDAG_BULLETS:
        bullet(doc, t)

    heading(doc, "Project Experience")
    role_line(doc, "Vision-Language Model (VLM) Service Integration", "Ongoing")
    italic_line(doc, "Python  |  REST APIs  |  Containerized services  |  Computer Vision and NLP")
    for t in VLM_BULLETS:
        bullet(doc, t)

    role_line(doc, "Python Automation — Multilingual Test-Case Translation Pipeline", "Self-directed")
    for t in AUTO_BULLETS:
        bullet(doc, t)

    role_line(
        doc,
        "Rental Property Analytics Tool  |  Data Science Intern, Saanvi Innovations Pvt Ltd",
        "2023",
    )
    for t in RENTAL_BULLETS:
        bullet(doc, t)

    role_line(
        doc,
        "Cervical Cancer Risk Prediction Model  |  Data Science Intern, Saanvi Innovations Pvt Ltd",
        "2023",
    )
    for t in CANCER_BULLETS:
        bullet(doc, t)

    heading(doc, "Education")
    role_line(
        doc,
        "Rustamji Institute of Technology, BSF Academy",
        "Gwalior, MP  |  Nov 2020 – Jun 2024",
    )
    body(
        doc,
        "Bachelor of Technology (B.Tech), Automobile Engineering  |  CGPA: 8.3 / 10  |  "
        "Rajiv Gandhi Proudyogiki Vishwavidyalaya, Bhopal",
        size=10,
        space_after=1,
    )

    heading(doc, "Certifications and Additional")
    bullet(doc, "Python Programming Essentials (Cisco); Data Science and Data Analytics (Kanishka IT Solutions); ANSYS and SolidWorks")
    bullet(
        doc,
        "Gen-AI Hackathon by EY (Participant); E-BAJA ATV Vehicle Challenge (Team Member); Coordinated Google Developer Fest 2023. Competencies: analytical problem solving, communication, teamwork, eagerness to learn cloud and AI.",
    )

    doc.save(path)


def pstyle(name, **kwargs):
    defaults = dict(
        fontName="Times-Roman",
        fontSize=9.5,
        leading=11.6,
        textColor=black,
        alignment=TA_LEFT,
    )
    defaults.update(kwargs)
    return ParagraphStyle(name, **defaults)


def build_pdf(path: Path):
    navy = HexColor("#1F4E79")
    styles = {
        "name": pstyle(
            "name",
            fontName="Times-Bold",
            fontSize=15,
            leading=17,
            alignment=TA_CENTER,
            textColor=navy,
        ),
        "contact": pstyle("contact", fontSize=9, leading=11, alignment=TA_CENTER),
        "h": pstyle(
            "h",
            fontName="Times-Bold",
            fontSize=10.5,
            leading=13,
            textColor=navy,
            spaceBefore=7,
            spaceAfter=1,
        ),
        "body": pstyle("body", fontSize=9.5, leading=11.8, alignment=TA_JUSTIFY, spaceAfter=1),
        "role": pstyle("role", fontName="Times-Bold", fontSize=10, leading=12, spaceBefore=4),
        "meta": pstyle(
            "meta",
            fontName="Times-Italic",
            fontSize=9,
            leading=11,
            textColor=HexColor("#333333"),
        ),
        "skill": pstyle("skill", fontSize=9.5, leading=11.6, spaceAfter=0.5),
        "bullet": pstyle("bullet", fontSize=9.5, leading=11.6, leftIndent=10, spaceAfter=1.2),
    }

    doc = SimpleDocTemplate(
        str(path),
        pagesize=letter,
        leftMargin=0.6 * inch,
        rightMargin=0.6 * inch,
        topMargin=0.4 * inch,
        bottomMargin=0.4 * inch,
    )
    story = []
    story.append(Paragraph(NAME, styles["name"]))
    story.append(Spacer(1, 3))
    story.append(Paragraph(CONTACT.replace("  |  ", " | "), styles["contact"]))

    def section(title):
        story.append(Paragraph(title.upper(), styles["h"]))
        story.append(
            HRFlowable(width="100%", thickness=0.8, color=navy, spaceBefore=0, spaceAfter=3)
        )

    def bullets(items):
        for t in items:
            story.append(Paragraph("•  " + t, styles["bullet"]))

    def job(left, right, italic=None):
        story.append(Paragraph(f"<b>{left}</b>  |  {right}", styles["role"]))
        if italic:
            story.append(Paragraph(italic, styles["meta"]))

    section("Professional Summary")
    story.append(Paragraph(SUMMARY, styles["body"]))

    section("Technical Skills")
    for label, items in SKILL_BLOCKS:
        story.append(Paragraph(f"<b>{label}:</b> {items}", styles["skill"]))

    section("Professional Experience")
    job(
        "EDAG Production Solutions India Pvt Ltd",
        "Pune, India | Jun 2024 – Present",
        "Android Automotive Developer | See-Through Surround-View Application",
    )
    bullets(EDAG_BULLETS)

    section("Project Experience")
    job(
        "Vision-Language Model (VLM) Service Integration",
        "Ongoing",
        "Python | REST APIs | Containerized services | Computer Vision and NLP",
    )
    bullets(VLM_BULLETS)

    job("Python Automation — Multilingual Test-Case Translation Pipeline", "Self-directed")
    bullets(AUTO_BULLETS)

    job(
        "Rental Property Analytics Tool | Data Science Intern, Saanvi Innovations Pvt Ltd",
        "2023",
    )
    bullets(RENTAL_BULLETS)

    job(
        "Cervical Cancer Risk Prediction Model | Data Science Intern, Saanvi Innovations Pvt Ltd",
        "2023",
    )
    bullets(CANCER_BULLETS)

    section("Education")
    job(
        "Rustamji Institute of Technology, BSF Academy",
        "Gwalior, MP | Nov 2020 – Jun 2024",
    )
    story.append(
        Paragraph(
            "Bachelor of Technology (B.Tech), Automobile Engineering | CGPA: 8.3 / 10 | "
            "Rajiv Gandhi Proudyogiki Vishwavidyalaya, Bhopal",
            styles["body"],
        )
    )

    section("Certifications and Additional")
    bullets(
        [
            "Python Programming Essentials (Cisco); Data Science and Data Analytics (Kanishka IT Solutions); ANSYS and SolidWorks",
            "Gen-AI Hackathon by EY (Participant); E-BAJA ATV Vehicle Challenge (Team Member); Coordinated Google Developer Fest 2023. Competencies: analytical problem solving, communication, teamwork, eagerness to learn cloud and AI.",
        ]
    )

    doc.build(story)


def build_txt(path: Path):
    lines = [
        NAME,
        CONTACT.replace("  |  ", " | "),
        "",
        "PROFESSIONAL SUMMARY",
        SUMMARY,
        "",
        "TECHNICAL SKILLS",
    ]
    for label, items in SKILL_BLOCKS:
        lines.append(f"{label}: {items}")
    lines += [
        "",
        "PROFESSIONAL EXPERIENCE",
        "",
        "EDAG Production Solutions India Pvt Ltd | Pune, India | Jun 2024 - Present",
        "Android Automotive Developer | See-Through Surround-View Application",
    ]
    lines += [f"- {t}" for t in EDAG_BULLETS]
    lines += [
        "",
        "PROJECT EXPERIENCE",
        "",
        "Vision-Language Model (VLM) Service Integration | Ongoing",
        "Python | REST APIs | Containerized services | Computer Vision and NLP",
    ]
    lines += [f"- {t}" for t in VLM_BULLETS]
    lines += [
        "",
        "Python Automation - Multilingual Test-Case Translation Pipeline | Self-directed",
    ]
    lines += [f"- {t}" for t in AUTO_BULLETS]
    lines += [
        "",
        "Rental Property Analytics Tool | Data Science Intern, Saanvi Innovations Pvt Ltd | 2023",
    ]
    lines += [f"- {t}" for t in RENTAL_BULLETS]
    lines += [
        "",
        "Cervical Cancer Risk Prediction Model | Data Science Intern, Saanvi Innovations Pvt Ltd | 2023",
    ]
    lines += [f"- {t}" for t in CANCER_BULLETS]
    lines += [
        "",
        "EDUCATION",
        "Rustamji Institute of Technology, BSF Academy | Gwalior, MP | Nov 2020 - Jun 2024",
        "Bachelor of Technology (B.Tech), Automobile Engineering | CGPA: 8.3 / 10 | Rajiv Gandhi Proudyogiki Vishwavidyalaya, Bhopal",
        "",
        "CERTIFICATIONS AND ADDITIONAL",
        "- Python Programming Essentials (Cisco); Data Science and Data Analytics (Kanishka IT Solutions); ANSYS and SolidWorks",
        "- Gen-AI Hackathon by EY (Participant); E-BAJA ATV Vehicle Challenge (Team Member); Coordinated Google Developer Fest 2023. Competencies: analytical problem solving, communication, teamwork, eagerness to learn cloud and AI.",
        "",
    ]
    path.write_text("\n".join(lines), encoding="utf-8")


if __name__ == "__main__":
    docx_path = OUT_DIR / "Pamella_Shome_Resume_Accenture.docx"
    pdf_path = OUT_DIR / "Pamella_Shome_Resume_Accenture.pdf"
    txt_path = OUT_DIR / "Pamella_Shome_Resume_Accenture.txt"
    build_docx(docx_path)
    build_pdf(pdf_path)
    build_txt(txt_path)
    print(f"Wrote {docx_path}")
    print(f"Wrote {pdf_path}")
    print(f"Wrote {txt_path}")
